import csv
from datetime import datetime, date

from django.contrib import messages
from django.contrib.auth import logout
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, FileResponse
from django.shortcuts import render
from django.urls import reverse_lazy
from django.views.generic import ListView
from django.views.generic.edit import CreateView, UpdateView, DeleteView
from docx import Document
from docx.shared import Pt, Inches

from home.dwaclasses import DwaCommon
from .forms import AddressCreateForm, AddressUpdateForm, ListByTypeForm
from .models import Address


@login_required
def start_page(request):
    return render(request, 'clc/start_page.html')


def login(request):
    return render(request, 'clc/start_page.html')


def logout_request(request):
    logout(request)
    messages.info(request, "Logged out")
    print(f"*** logout **** request: {request}")
    return render(request, 'clc/start_page.html')


def export_envelope_list(request):
    print(f"\n*** export_envelope_list ***")
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="envelope_list.csv"'

    writer = csv.writer(response)
    columns = ['envelopeno', 'firstname', 'lastname', 'addrline1', 'addrline2', 'city', 'state', 'postalcode',
               'address_type']
    writer.writerow(columns)
    members = Address.objects.all().values_list(*columns).filter(envelopeno__gte=0, envelopeno__lte=999).order_by(
        'lastname')
    for member in members:
        writer.writerow(member)
    return response


def export_members_csv(request):
    print(f"\n*** export_members_csv ***")
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="members.csv"'

    writer = csv.writer(response)
    columns = ['firstname', 'lastname', 'addrline1', 'addrline2', 'city', 'state', 'postalcode', 'address_type']
    writer.writerow(columns)

    members = Address.objects.all().values_list(*columns).filter(address_type='m').order_by('lastname')
    for member in members:
        writer.writerow(member)
    return response


def export_advent_list(request):
    print(f"\n*** export_advent_csv ***")
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="advent.csv"'

    writer = csv.writer(response)
    columns = ['id', 'advent', 'firstname', 'lastname', 'addrline1', 'addrline2', 'city', 'state', 'postalcode',
               'address_type']
    writer.writerow(columns)

    members = Address.objects.all().values_list(*columns).filter(archive='N').order_by('advent', 'lastname')
    for member in members:
        writer.writerow(member)
    return response


def create_directory(self):
    print("\n ** entering create_directory")
    document = Document()
    sections = document.sections
    for section in sections:
        section.top_margin = Inches(0.25)
        section.bottom_margin = Inches(0.25)
        section.left_margin = Inches(0.50)
        section.right_margin = Inches(0.25)
        section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    run = paragraph.add_run("Christ's Lutheran Church Directory")
    font = run.font
    font.size = Pt(24)
    footer = section.footer
    paragraph = footer.paragraphs[0]
    timenow = datetime.now()
    date_string = timenow.strftime("%b %d, %Y")
    datenow = date.today()
    paragraph.text = f"Created on {date_string}"
    addresses = Address.objects.all().order_by('lastname').filter(archive='N', directorytype='d')
    row_index = 0
    cell_index = 0
    max_rows = 7

    add_row = True
    new_table = True

    for address in addresses:
        if new_table:
            document.add_paragraph()
            table = document.add_table(rows=0, cols=3)
            table.style = 'Table Grid'
            new_table = False
        if add_row:
            table.add_row()
            add_row = False
        print(f"address: {address.lastname}")

        cells = table.rows[row_index].cells
        content = address.format_cell()
        #        print (f"\n create_directory: {row_index},{cell_index}; content: {content}")
        cells[cell_index].text = f"{content}\n"
        cells[cell_index].width = Inches(2.5)
        cell_index += 1
        if cell_index == 3:
            #            print (f"add row {row_index}")
            add_row = True
            row_index += 1
            cell_index = 0
        if row_index == max_rows:
            row_index = 0
            document.add_page_break()
            new_table = True

    document.save("directory.docx")
    return FileResponse(open('directory.docx', 'rb'), as_attachment=True,
                        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


class ListByType(ListView):
    model = Address
    form_class = ListByTypeForm
    template_name = 'clc/list_by_type.html'
    type_filter = None

    def get(self, request, *args, **kwargs):
        print(f"\n*** ListByType::get {request.GET}")
        print(f"*** ListByType::session {request.session}; type_filter: {request.session.get('type_filter')}")
        self.type_filter = request.GET.get('address_type')
        if self.type_filter is None:
            self.type_filter = request.session.get('type_filter', 'm')
            print(f"*** ListByType::get session type filter {self.type_filter}")
            if self.type_filter is None:
                self.type_filter = 'm'
        print(f"type_filter: {self.type_filter}")
        request.session['type_filter'] = self.type_filter
        self.object_list = self.get_queryset()
        context = self.get_context_data()
        context['type_filter'] = self.type_filter

        return super().get(request, *args, **kwargs)

    def get_queryset(self):
        print(f"\n*** ListByType::get_queryset type: {self.type_filter}")
        if self.type_filter == "*":
            addresses = Address.objects.all().order_by('archive', 'lastname')
        else:
            addresses = Address.objects.all().order_by('archive', 'address_type', 'lastname').filter(
                address_type=self.type_filter)
        return addresses


class IndexView(DwaCommon, ListView):

    def get(self, request, *args, **kwargs):
        return self.dwa_get_limited(request, *args, **kwargs)

    def get_queryset(self):
        """Return the address objects."""
        return Address.objects.all().order_by('archive', 'lastname')


class AddressCreate(DwaCommon, CreateView):
    model = Address
    form_class = AddressCreateForm

    def form_valid(self, form):
        print("*** AddrCreate:form_valid enter\n")
        return self.dwa_form_valid_save(form)


class AddressUpdate(DwaCommon, UpdateView):
    model = Address
    form_class = AddressUpdateForm

    def form_valid(self, form):
        print("*** AddressUpdate:form_valid enter\n")
        return self.dwa_form_valid_save(form)


class AddressDelete(DeleteView):
    model = Address
    success_url = reverse_lazy('clc:address_list')
